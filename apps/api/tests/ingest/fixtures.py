"""Sample course files built in memory, for ingestion tests.

Nothing binary is checked in: each builder writes a small, realistic file with
pymupdf, python-pptx, python-docx or Pillow, so a test (here, or in the
chunking, matching and pipeline tests) can see exactly what is inside it.

What each file holds:

``calculus_pdf``   p. 1 "The Chain Rule" (with :data:`CHAIN_EXAMPLE`), p. 2 "The Product Rule",
                   p. 3 a picture of handwritten-style work with no text layer.
``calculus_pptx``  slide 1 and 2 titled "The chain rule" (bullets, then a table), slide 3
                   "Product rule" (with a typed equation), slide 4 a picture and nothing else.
``calculus_docx``  an intro paragraph, then Heading 1 "The chain rule" (with a typed equation
                   and a table), Heading 2 "Worked example", Heading 1 "The product rule".
``chain_rule_png`` :data:`PNG_TEXT` typed large enough for Tesseract.
"""

from __future__ import annotations

import io
import struct
import zipfile

import pymupdf
from PIL import Image, ImageDraw, ImageFont

CHAIN_EXAMPLE = "d/dx sin(x^3+1) = cos(x^3+1) * 3x^2"
PRODUCT_EXAMPLE = "d/dx x^2 sin(x) = 2x sin(x) + x^2 cos(x)"
PNG_TEXT = "Chain rule: dy/dx = dy/du * du/dx"
HANDWRITTEN_LINES = ["u = x^3 + 1", "du/dx = 3x^2"]
PDF_TITLE = "Week 3: Differentiation rules"
DOCX_TITLE = "Chain rule notes"

_M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
_A14 = "http://schemas.microsoft.com/office/drawing/2010/main"


# ---------------------------------------------------------------------------
# Images


def text_image(text: str | list[str], *, size: int = 48, margin: int = 40) -> Image.Image:
    """Black text on white, one line per item, in Pillow's bundled font."""
    lines = [text] if isinstance(text, str) else list(text)
    font = ImageFont.load_default(size=size)
    measure = ImageDraw.Draw(Image.new("L", (1, 1)))
    width = max(measure.textbbox((0, 0), line, font=font)[2] for line in lines) + 2 * margin
    line_height = int(size * 1.5)
    img = Image.new("RGB", (width, line_height * len(lines) + 2 * margin), "white")
    draw = ImageDraw.Draw(img)
    for i, line in enumerate(lines):
        draw.text((margin, margin + i * line_height), line, fill="black", font=font)
    return img


def _encode(img: Image.Image, fmt: str, **options) -> bytes:
    out = io.BytesIO()
    img.save(out, format=fmt, **options)
    return out.getvalue()


def text_png(text: str | list[str], *, size: int = 48) -> bytes:
    return _encode(text_image(text, size=size), "PNG")


def chain_rule_png() -> bytes:
    """:data:`PNG_TEXT` at 48 px, which Tesseract reads reliably."""
    return text_png(PNG_TEXT)


def chain_rule_jpeg() -> bytes:
    return _encode(text_image(PNG_TEXT), "JPEG", quality=90)


def chain_rule_webp() -> bytes:
    return _encode(text_image(PNG_TEXT), "WEBP", quality=90)


def photo_jpeg(width: int = 4000, height: int = 3000, *, orientation: int = 6) -> bytes:
    """A large phone-style photo whose EXIF says to rotate it 90 degrees (orientation 6)."""
    img = Image.new("RGB", (width, height), (235, 232, 220))
    ImageDraw.Draw(img).text((200, 200), PNG_TEXT, fill="black", font=ImageFont.load_default(size=160))
    exif = Image.Exif()
    exif[0x0112] = orientation
    return _encode(img, "JPEG", quality=80, exif=exif)


# ---------------------------------------------------------------------------
# PDF


def _write_lines(page: pymupdf.Page, top: float, lines: list[str], size: float = 11) -> None:
    page.insert_text((72, top), "\n".join(lines), fontsize=size, fontname="helv", lineheight=1.4)


def calculus_pdf(*, title: str | None = PDF_TITLE) -> bytes:
    """Three pages: chain rule, product rule, and one with only a picture."""
    doc = pymupdf.open()

    page = doc.new_page()
    page.insert_text((72, 90), "The Chain Rule", fontsize=22, fontname="hebo")
    _write_lines(page, 130, [
        "When one function is inside another, differentiate the outside function",
        "and multiply by the derivative of the inside function.",
        "If y = f(g(x)), then dy/dx = f'(g(x)) g'(x).",
        f"Example: {CHAIN_EXAMPLE}",
        "Here the inside function is u = x^3 + 1 and the outside function is sin(u).",
    ])
    page.insert_text((300, 760), "1", fontsize=9, fontname="helv")

    page = doc.new_page()
    page.insert_text((72, 90), "The Product Rule", fontsize=22, fontname="hebo")
    _write_lines(page, 130, [
        "To differentiate a product, differentiate each factor in turn and add.",
        "d/dx [f(x) g(x)] = f'(x) g(x) + f(x) g'(x)",
        f"Example: {PRODUCT_EXAMPLE}",
    ])
    page.insert_text((300, 760), "2", fontsize=9, fontname="helv")

    page = doc.new_page()
    page.insert_image(pymupdf.Rect(60, 80, 552, 400), stream=text_png(HANDWRITTEN_LINES, size=64))

    if title is not None:
        doc.set_metadata({"title": title})
    data = doc.tobytes()
    doc.close()
    return data


def encrypted_pdf(password: str = "calculus") -> bytes:
    """:func:`calculus_pdf` locked with a user password."""
    doc = pymupdf.open(stream=calculus_pdf(), filetype="pdf")
    data = doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw=password, owner_pw=password + "-owner")
    doc.close()
    return data


def damaged_pdf() -> bytes:
    """A PDF header followed by bytes that are not a PDF."""
    return b"%PDF-1.7\n" + bytes(range(256)) * 8


def blank_pdf(pages: int = 2) -> bytes:
    doc = pymupdf.open()
    for _ in range(pages):
        doc.new_page()
    data = doc.tobytes()
    doc.close()
    return data


def long_pdf(pages: int) -> bytes:
    doc = pymupdf.open()
    for i in range(pages):
        doc.new_page(width=300, height=300).insert_text(
            (20, 40), f"Page {i + 1}: the derivative of x^{i + 2} is {i + 2}x^{i + 1}.", fontsize=9)
    data = doc.tobytes()
    doc.close()
    return data


# ---------------------------------------------------------------------------
# Office Math (shared by the PowerPoint and Word builders)


def _mr(text: str) -> str:
    return f'<m:r><m:t xml:space="preserve">{text}</m:t></m:r>'


def _mf(num: str, den: str) -> str:
    return f"<m:f><m:num>{num}</m:num><m:den>{den}</m:den></m:f>"


CHAIN_OMML = _mf(_mr("dy"), _mr("dx")) + _mr("=") + _mf(_mr("dy"), _mr("du")) + _mr("·") + _mf(_mr("du"), _mr("dx"))
"""dy/dx = dy/du · du/dx as Office Math."""
CHAIN_LINEAR = "dy/dx = dy/du · du/dx"

PRODUCT_OMML = (_mf(_mr("d"), _mr("dx")) + f"<m:d><m:e>{_mr('uv')}</m:e></m:d>" + _mr(" = ") + _mr("u")
                + _mf(_mr("dv"), _mr("dx")) + _mr(" + ") + _mr("v") + _mf(_mr("du"), _mr("dx")))
"""d/dx (uv) = u dv/dx + v du/dx as Office Math."""
PRODUCT_LINEAR = "d/dx (uv) = u dv/dx + v du/dx"


# ---------------------------------------------------------------------------
# PowerPoint


def _equation_box(slide, omml: str, left: int, top: int, width: int, height: int) -> None:
    """A text box holding an equation, wrapped in mc:AlternateContent the way PowerPoint saves it."""
    from lxml import etree
    from pptx.oxml.ns import qn

    sp = slide.shapes.add_textbox(left, top, width, height)._element
    paragraph = sp.txBody.find(qn("a:p"))
    paragraph.append(etree.fromstring(
        f'<a14:m xmlns:a14="{_A14}" xmlns:m="{_M}"><m:oMathPara><m:oMath>{omml}</m:oMath></m:oMathPara></a14:m>'))
    wrapper = etree.fromstring(
        f'<mc:AlternateContent xmlns:mc="{_MC}"><mc:Choice xmlns:a14="{_A14}" Requires="a14"/>'
        f'<mc:Fallback/></mc:AlternateContent>')
    sp.addprevious(wrapper)
    wrapper[0].append(sp)


def calculus_pptx() -> bytes:
    """Four slides: two titled "The chain rule", one "Product rule", one picture only."""
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    title_and_content, title_only, blank = deck.slide_layouts[1], deck.slide_layouts[5], deck.slide_layouts[6]

    slide = deck.slides.add_slide(title_and_content)
    slide.shapes.title.text = "The chain rule"
    body = slide.placeholders[1].text_frame
    body.text = "If y = f(g(x)), then dy/dx = f'(g(x)) g'(x)."
    body.add_paragraph().text = f"Example: {CHAIN_EXAMPLE}"
    slide.notes_slide.notes_text_frame.text = "Remind students to multiply by the derivative of the inside."

    slide = deck.slides.add_slide(title_only)
    slide.shapes.title.text = "The chain rule"
    table = slide.shapes.add_table(3, 3, Inches(0.5), Inches(2), Inches(9), Inches(2)).table
    for r, row in enumerate([("Outer", "Inner", "Derivative"),
                             ("sin(u)", "u = x^3 + 1", "cos(x^3+1) * 3x^2"),
                             ("e^u", "u = 2x", "2e^(2x)")]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value

    slide = deck.slides.add_slide(title_and_content)
    slide.shapes.title.text = "Product rule"
    body = slide.placeholders[1].text_frame
    body.text = "(fg)' = f'g + fg'"
    body.add_paragraph().text = f"Example: {PRODUCT_EXAMPLE}"
    _equation_box(slide, PRODUCT_OMML, Inches(1), Inches(6), Inches(8), Inches(1))

    slide = deck.slides.add_slide(blank)
    picture = io.BytesIO(_encode(text_image(HANDWRITTEN_LINES, size=64), "JPEG", quality=90))
    slide.shapes.add_picture(picture, Inches(0.5), Inches(0.5), width=Inches(9))

    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()


def long_pptx(slides: int) -> bytes:
    from pptx import Presentation

    deck = Presentation()
    for i in range(slides):
        deck.slides.add_slide(deck.slide_layouts[5]).shapes.title.text = f"Slide {i + 1}"
    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()


# ---------------------------------------------------------------------------
# Word


def calculus_docx(*, title: str | None = DOCX_TITLE, with_picture: bool = False) -> bytes:
    """An intro, then Heading sections; ``with_picture`` adds a section that is only a picture."""
    import docx
    from docx.oxml import parse_xml

    document = docx.Document()
    if title is not None:
        document.core_properties.title = title
    document.add_paragraph("Notes for week 3 of Calculus I. Read these before the problem set.")

    document.add_heading("The chain rule", level=1)
    document.add_paragraph("If y = f(g(x)), differentiate the outside function and multiply by the derivative "
                           "of the inside function.")
    paragraph = document.add_paragraph("In Leibniz notation: ")
    paragraph._p.append(parse_xml(f'<m:oMath xmlns:m="{_M}">{CHAIN_OMML}</m:oMath>'))
    document.add_paragraph(f"Example: {CHAIN_EXAMPLE}")
    table = document.add_table(rows=3, cols=2)
    for r, row in enumerate([("Function", "Derivative"), ("sin(x^3+1)", "cos(x^3+1) * 3x^2"),
                             ("e^(2x)", "2e^(2x)")]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value

    document.add_heading("Worked example", level=2)
    document.add_paragraph("Find the derivative of (x^2+1)^5. The inside is u = x^2 + 1, so the answer is "
                           "5(x^2+1)^4 * 2x.")

    document.add_heading("The product rule", level=1)
    document.add_paragraph("d/dx [f(x) g(x)] = f'(x) g(x) + f(x) g'(x)")
    document.add_paragraph(f"Example: {PRODUCT_EXAMPLE}")

    if with_picture:
        document.add_heading("Handwritten example", level=1)
        document.add_picture(io.BytesIO(text_png(HANDWRITTEN_LINES, size=64)))

    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def empty_docx() -> bytes:
    import docx

    out = io.BytesIO()
    docx.Document().save(out)
    return out.getvalue()


# ---------------------------------------------------------------------------
# Hostile and unsupported files


def zip_bomb_docx(declared_bytes: int = 300 * 1024 * 1024) -> bytes:
    """A tiny docx-shaped zip whose central directory claims one member is huge."""
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "<document/>")
        archive.writestr("word/media/huge.bin", b"\0" * 1024)
    data = bytearray(out.getvalue())
    # Central directory header: the uncompressed size sits 24 bytes in, the name 46 bytes in.
    header = data.rfind(b"word/media/huge.bin") - 46
    assert data[header:header + 4] == b"PK\x01\x02"
    struct.pack_into("<I", data, header + 24, declared_bytes)
    return bytes(data)


def encrypted_office() -> bytes:
    """A stand-in for a password-protected .docx: an OLE container naming an EncryptedPackage stream."""
    return (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 504 + "EncryptedPackage".encode("utf-16-le")
            + b"\0" * 512)


def plain_zip() -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("notes.txt", "not an office file")
    return out.getvalue()
