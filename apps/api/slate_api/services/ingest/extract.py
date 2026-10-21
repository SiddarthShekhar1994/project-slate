"""Turn an uploaded file into pages of text, the first step of ingestion.

Uploads come from students, so nothing about them is trusted: the type is
decided from the file's first bytes rather than its name, archives are checked
for zip bombs before anything unpacks them, and every limit (pages, pixels,
image bytes) is fixed so one huge upload cannot exhaust the worker.

Each format keeps the unit a student would cite: a PDF page ("p. 3"), a slide
("slide 7"), a Word section under one heading ("section 2"), or the whole
image. ``Page.text`` holds everything readable on that unit, heading included.
Pages whose text layer is missing or too thin (scans, photos of handwriting,
slides that are one big picture) are marked ``needs_ocr`` and carry a PNG for
:mod:`.ocr`. Equations typed in Word or PowerPoint are kept as linear text
(``dy/dx = dy/du*du/dx``) so the notation detector and skill matcher still
see them.
"""

from __future__ import annotations

import io
import math
import re
import zipfile
from collections import Counter
from collections.abc import Callable, Iterator
from dataclasses import dataclass, field
from pathlib import PureWindowsPath
from typing import Any

import pymupdf
from PIL import Image, ImageOps

from .types import Extracted, Page

PDF = "application/pdf"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PNG = "image/png"
JPEG = "image/jpeg"
WEBP = "image/webp"
SUPPORTED_MIMES = (PDF, PPTX, DOCX, PNG, JPEG, WEBP)

MAX_PAGES = 300
"""Pages, slides or sections read from one file. The rest is dropped with a warning."""
MAX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
MAX_ZIP_ENTRIES = 20_000
OCR_DPI = 150
MAX_RENDER_PIXELS = 4_000_000
MAX_IMAGE_SIDE = 2000
MAX_SOURCE_PIXELS = 80_000_000
"""Images bigger than this are refused before decoding (a 12 MP phone photo is 12 million)."""
MAX_OCR_IMAGE_BYTES = 150 * 1024 * 1024
"""PNGs held for OCR across one file. A long scanned PDF stops getting images past this."""

MIN_TEXT_CHARS = 40
"""A page with fewer non-space characters than this has no usable text layer."""
MOSTLY_IMAGE = 0.6
MOSTLY_IMAGE_MAX_CHARS = 200
"""A page this much covered by pictures, with little text, is read with OCR as well."""

HEADING_RATIO = 1.2
HEADING_TOP = 0.35
MAX_HEADING_CHARS = 200

UNSUPPORTED = ("Slate can read PDFs, PowerPoint slides (.pptx), Word documents (.docx) and PNG, JPEG or WEBP "
               "images. This file isn't one of those.")
EMPTY = "This file is empty."
TOO_BIG_ARCHIVE = ("This file unpacks to more than 200 MB, which is more than Slate can read. "
                   "Try removing videos or large pictures, or splitting it into smaller files.")
OFFICE_PASSWORD = "This file is password protected. Remove the password, save it again and upload the new copy."
OFFICE_LEGACY = "This looks like an older .ppt or .doc file. Save it as .pptx or .docx and upload it again."
PDF_PASSWORD = "This PDF is password protected. Remove the password, save it again and upload the new copy."
PDF_DAMAGED = "This PDF looks damaged and couldn't be opened. Try exporting it again and uploading the new copy."
PDF_EMPTY = "This PDF has no pages."
PDF_BLANK = "This PDF doesn't have any text or pictures in it."
PPTX_DAMAGED = ("This presentation looks damaged and couldn't be opened. Try saving it again from PowerPoint, "
                "Keynote or Google Slides and uploading the new copy.")
PPTX_EMPTY = "This presentation has no slides."
DOCX_DAMAGED = "This document looks damaged and couldn't be opened. Try saving it again and uploading the new copy."
DOCX_EMPTY = "This document doesn't have any text or pictures in it."
IMAGE_DAMAGED = "This image looks damaged and couldn't be opened. Try taking or exporting it again."
IMAGE_TOO_BIG = "This image is too large to read. Try a smaller photo or a screenshot."
IMAGE_BUDGET = ("Some pages are scans or pictures that Slate didn't have room to read. Splitting the file into "
                "smaller parts will let Slate read all of them.")

_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_ENCRYPTED_PACKAGE = "EncryptedPackage".encode("utf-16-le")

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
_SLIDE_SHAPE_TAGS = {f"{{{_P}}}{name}" for name in ("sp", "grpSp", "graphicFrame", "cxnSp", "pic")}


class ExtractionError(Exception):
    """The file can't be read. ``message`` is written for the student."""

    def __init__(self, message: str):
        super().__init__(message)
        self.message = message


@dataclass
class ExtractedFile(Extracted):
    """:class:`Extracted` plus what was left out, so the library can tell the student.

    ``page_count`` is the number of pages, slides or sections in the whole
    file; ``pages`` stops at :data:`MAX_PAGES`, and then ``truncated`` is set
    and ``warnings`` says so in a sentence.
    """

    truncated: bool = False
    warnings: list[str] = field(default_factory=list)


# ---------------------------------------------------------------------------
# Type detection


def detect_mime(data: bytes, filename: str = "") -> str | None:
    """The file's MIME type from its first bytes, or ``None`` if Slate can't read it.

    ``filename`` is never trusted: a PNG renamed to ``.pdf`` is a PNG. Office
    files are zip archives, which are opened (without unpacking) to tell slides
    from documents and to refuse ones whose declared size is over 200 MB.
    """
    if not data:
        return None
    if data.startswith(b"\x89PNG\r\n\x1a\n"):
        return PNG
    if data.startswith(b"\xff\xd8\xff"):
        return JPEG
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return WEBP
    if data.startswith(b"PK\x03\x04"):
        mime, _problem = _inspect_zip(data)
        return mime
    # Acrobat accepts a few bytes of junk before the header, and so does MuPDF.
    if b"%PDF-" in data[:1024]:
        return PDF
    return None


def rejection_message(data: bytes, filename: str = "") -> str:
    """A sentence for the student explaining why :func:`detect_mime` returned ``None``."""
    if not data:
        return EMPTY
    if data.startswith(_OLE_MAGIC):
        return OFFICE_PASSWORD if _ENCRYPTED_PACKAGE in data else OFFICE_LEGACY
    if data.startswith(b"PK\x03\x04"):
        _mime, problem = _inspect_zip(data)
        if problem:
            return problem
    return UNSUPPORTED


def _inspect_zip(data: bytes) -> tuple[str | None, str | None]:
    """(mime, problem) for a zip archive, reading only its central directory."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
    except (zipfile.BadZipFile, zipfile.LargeZipFile, ValueError, OSError, EOFError, NotImplementedError):
        return None, None
    # zipfile stops reading a member at its declared size, so bounding the declared
    # total also bounds what python-pptx and python-docx can unpack later.
    if len(infos) > MAX_ZIP_ENTRIES or sum(info.file_size for info in infos) > MAX_UNCOMPRESSED_BYTES:
        return None, TOO_BIG_ARCHIVE
    names = [info.filename for info in infos]
    if "[Content_Types].xml" not in names:
        return None, None
    if any(name.startswith("ppt/") for name in names):
        return PPTX, None
    if any(name.startswith("word/") for name in names):
        return DOCX, None
    return None, None


# ---------------------------------------------------------------------------
# Entry point


def extract(data: bytes, mime: str, filename: str = "") -> Extracted:
    """Split a file into :class:`Page` objects. Returns an :class:`ExtractedFile`.

    Raises :class:`ExtractionError` with a student-friendly message when the
    file is empty, encrypted, damaged or not a supported type.
    """
    if not data:
        raise ExtractionError(EMPTY)
    readers: dict[str, Callable[[bytes, str], ExtractedFile]] = {
        PDF: _extract_pdf, PPTX: _extract_pptx, DOCX: _extract_docx,
        PNG: _extract_image, JPEG: _extract_image, WEBP: _extract_image,
    }
    reader = readers.get(mime)
    if reader is None:
        raise ExtractionError(UNSUPPORTED)
    return reader(data, filename)


def _filename_title(filename: str) -> str:
    # PureWindowsPath splits on both / and \, so browser-supplied paths lose their folders.
    stem = PureWindowsPath(filename or "").stem.strip() if filename else ""
    return _one_line(stem) or "Untitled"


def _truncation_warning(total: int, unit: str) -> str:
    return f"This file has {total} {unit}s. Slate read the first {MAX_PAGES}."


class _ImageBudget:
    """Caps the PNG bytes kept for OCR across one file."""

    def __init__(self, limit: int = MAX_OCR_IMAGE_BYTES):
        self.remaining = limit
        self.exhausted = False

    def take(self, png: bytes | None) -> bytes | None:
        if png is None:
            return None
        if len(png) > self.remaining:
            self.exhausted = True
            return None
        self.remaining -= len(png)
        return png


def _finish(result: ExtractedFile, budget: _ImageBudget) -> ExtractedFile:
    if budget.exhausted:
        result.warnings.append(IMAGE_BUDGET)
    return result


# ---------------------------------------------------------------------------
# Text helpers

_CONTROL = re.compile(r"[\x00-\x08\x0c\x0e-\x1f\x7f]")


def _clean(text: str) -> str:
    """Normalize line endings and drop control characters (Postgres rejects NUL)."""
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x0b", "\n")
    text = _CONTROL.sub("", text)
    lines = [re.sub(r" {2,}", " ", line).strip(" ").rstrip() for line in text.split("\n")]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def _one_line(text: str) -> str:
    return re.sub(r"\s+", " ", _CONTROL.sub("", text)).strip()[:MAX_HEADING_CHARS]


def _content_chars(text: str) -> int:
    return len(re.sub(r"\s", "", text))


def _needs_ocr(text: str, picture_share: float, has_picture: bool) -> bool:
    chars = _content_chars(text)
    if chars < MIN_TEXT_CHARS:
        return has_picture
    return picture_share >= MOSTLY_IMAGE and chars < MOSTLY_IMAGE_MAX_CHARS


# ---------------------------------------------------------------------------
# Images


class _ImageTooLarge(ValueError):
    pass


def _open_image(data: bytes) -> Image.Image:
    """Decode an image upright, flattened onto white, and at most MAX_IMAGE_SIDE on the long side.

    Raises _ImageTooLarge when the image is too large to decode safely, and
    whatever Pillow raises when it is damaged.
    """
    img = Image.open(io.BytesIO(data))
    if img.width * img.height > MAX_SOURCE_PIXELS:
        raise _ImageTooLarge(f"{img.width}x{img.height}")
    if img.format == "JPEG":
        # Decode big JPEGs at a reduced scale instead of at full size.
        img.draft("RGB", (MAX_IMAGE_SIDE, MAX_IMAGE_SIDE))
    img = ImageOps.exif_transpose(img)
    if img.mode in ("RGBA", "LA", "PA") or (img.mode == "P" and "transparency" in img.info):
        rgba = img.convert("RGBA")
        img = Image.new("RGB", rgba.size, "white")
        img.paste(rgba, mask=rgba.getchannel("A"))
    elif img.mode not in ("RGB", "L"):
        img = img.convert("RGB")
    img.thumbnail((MAX_IMAGE_SIDE, MAX_IMAGE_SIDE), Image.Resampling.LANCZOS)
    return img


def _png_bytes(img: Image.Image) -> bytes:
    out = io.BytesIO()
    img.save(out, format="PNG")
    return out.getvalue()


def _picture_png(blob: bytes) -> bytes | None:
    """An embedded picture as PNG, or None for formats Pillow can't draw (EMF, WMF, SVG)."""
    try:
        return _png_bytes(_open_image(blob))
    except Exception:  # Pillow raises many kinds of error for formats and files it can't decode
        return None


def _extract_image(data: bytes, filename: str) -> ExtractedFile:
    try:
        png = _png_bytes(_open_image(data))
    except (_ImageTooLarge, Image.DecompressionBombError):
        raise ExtractionError(IMAGE_TOO_BIG) from None
    except Exception as exc:  # Pillow raises many kinds of error for damaged files
        raise ExtractionError(IMAGE_DAMAGED) from exc
    page = Page(number=1, label="image", needs_ocr=True, image=png)
    return ExtractedFile(pages=[page], title=_filename_title(filename), page_count=1)


# ---------------------------------------------------------------------------
# PDF

_PDF_TEXT_FLAGS = (pymupdf.TEXTFLAGS_DICT & ~pymupdf.TEXT_PRESERVE_IMAGES & ~pymupdf.TEXT_PRESERVE_LIGATURES)


@dataclass
class _Line:
    text: str
    size: float
    top: float
    block: int


def _extract_pdf(data: bytes, filename: str) -> ExtractedFile:
    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:  # MuPDF raises several error types for damaged files
        raise ExtractionError(PDF_DAMAGED) from exc
    with doc:
        if doc.needs_pass:
            raise ExtractionError(PDF_PASSWORD)
        total = doc.page_count
        if total == 0:
            raise ExtractionError(PDF_EMPTY)
        budget = _ImageBudget()
        pages: list[Page] = []
        warnings: list[str] = []
        for index in range(min(total, MAX_PAGES)):
            try:
                pages.append(_pdf_page(doc[index], index + 1, budget))
            except Exception:  # one broken page should not lose the rest of the file
                pages.append(Page(number=index + 1, label=f"p. {index + 1}"))
                warnings.append(f"Slate couldn't read page {index + 1}.")
        if all(not p.text and not p.needs_ocr for p in pages):
            raise ExtractionError(PDF_DAMAGED if warnings else PDF_BLANK)
        title = _pdf_title(doc.metadata or {}) or _filename_title(filename)
    if total > MAX_PAGES:
        warnings.insert(0, _truncation_warning(total, "page"))
    result = ExtractedFile(pages=pages, title=title, page_count=total, truncated=total > MAX_PAGES,
                           warnings=warnings)
    return _finish(result, budget)


def _pdf_title(metadata: dict[str, Any]) -> str:
    title = _one_line(str(metadata.get("title") or ""))
    # Office exports often stamp the app name or a placeholder in the title.
    title = re.sub(r"^Microsoft (Word|PowerPoint) - ", "", title)
    if title.lower() in {"untitled", "untitled document", "powerpoint presentation", "slide 1"}:
        return ""
    return title


def _pdf_page(page: pymupdf.Page, number: int, budget: _ImageBudget) -> Page:
    raw = page.get_text("dict", flags=_PDF_TEXT_FLAGS, sort=True)
    lines: list[_Line] = []
    blocks: list[str] = []
    for block_no, block in enumerate(raw.get("blocks", [])):
        if block.get("type") != 0:
            continue
        block_lines = []
        for line in block.get("lines", []):
            spans = line.get("spans", [])
            text = "".join(span.get("text", "") for span in spans).strip()
            if not text:
                continue
            size = max(span.get("size", 0.0) for span in spans if span.get("text", "").strip())
            lines.append(_Line(text, size, line["bbox"][1], block_no))
            block_lines.append(text)
        if block_lines:
            blocks.append("\n".join(block_lines))
    text = _clean("\n\n".join(blocks))
    heading = _pdf_heading(lines, page.rect)

    # Handwriting exported from a tablet app is vector strokes, not a picture, so a
    # thin text layer always counts; a render that comes out one flat color is a blank page.
    needs_ocr = _needs_ocr(text, _picture_share(page), has_picture=True)
    image = None
    if needs_ocr:
        pix = _render(page)
        if pix.is_unicolor:
            needs_ocr = False
        else:
            image = budget.take(pix.tobytes("png"))
    return Page(number=number, label=f"p. {number}", heading=heading, text=text, needs_ocr=needs_ocr, image=image)


def _pdf_heading(lines: list[_Line], rect: pymupdf.Rect) -> str:
    """The largest line near the top, if it is clearly bigger than the body text."""
    if not lines:
        return ""
    weights: Counter[float] = Counter()
    for line in lines:
        weights[round(line.size * 2) / 2] += len(line.text)
    top_zone = rect.y0 + rect.height * HEADING_TOP
    candidates = [i for i, line in enumerate(lines) if line.top <= top_zone]
    if not candidates:
        return ""
    best = max(candidates, key=lambda i: (lines[i].size, -i))
    head = lines[best]
    # The body size is measured without the heading itself, so a short page still has one.
    body_weights = weights.copy()
    body_weights[round(head.size * 2) / 2] -= sum(len(l.text) for l in lines
                                                    if abs(l.size - head.size) < 0.5 and l.top <= top_zone)
    body = next((size for size, count in body_weights.most_common() if count > 0), None)
    if body is None or head.size < body * HEADING_RATIO:
        return ""
    parts = [head.text]
    for line in lines[best + 1: best + 3]:  # a heading that wraps onto a second or third line
        if line.block != head.block or abs(line.size - head.size) >= 0.5:
            break
        parts.append(line.text)
    return _one_line(" ".join(parts))


def _picture_share(page: pymupdf.Page) -> float:
    area = abs(page.rect)
    if not area:
        return 0.0
    covered = 0.0
    for info in page.get_image_info():
        box = pymupdf.Rect(info["bbox"]) & page.rect
        if not box.is_empty:
            covered += abs(box)
    return min(covered / area, 1.0)


def _render(page: pymupdf.Page) -> pymupdf.Pixmap:
    zoom = OCR_DPI / 72
    pixels = page.rect.width * zoom * page.rect.height * zoom
    if pixels > MAX_RENDER_PIXELS:
        zoom *= math.sqrt(MAX_RENDER_PIXELS / pixels)
    return page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=False)


# ---------------------------------------------------------------------------
# Office XML text, shared by slides and documents


def _local(tag: Any) -> tuple[str, str]:
    """(namespace, local name) of an lxml tag; comments and processing instructions get ("", "")."""
    if not isinstance(tag, str):
        return "", ""
    if tag.startswith("{"):
        ns, _, name = tag[1:].partition("}")
        return ns, name
    return "", tag


class _TextWalker:
    """Collects the visible text under an Office XML element, in document order.

    Equations become linear text, deleted revisions and field codes are
    skipped, and the fallback copy inside ``mc:AlternateContent`` is ignored
    so text boxes are not read twice. Embedded picture ids are recorded too.
    """

    def __init__(self) -> None:
        self.parts: list[str] = []
        self.picture_ids: list[str] = []

    def text(self) -> str:
        return "".join(self.parts)

    def walk(self, el: Any) -> _TextWalker:
        ns, name = _local(el.tag)
        if ns == _M and name in ("oMath", "oMathPara"):
            self.parts.append(" " + omml_text(el).strip() + " ")
            return self
        if (ns == _MC and name == "Fallback") or (ns == _W and name in ("del", "delText", "instrText", "moveFrom")):
            return self
        if ns in (_W, _A) and name.endswith("Pr"):
            return self
        if ns in (_W, _A) and name == "t":
            self.parts.append(el.text or "")
        elif ns == _W and name == "tab":
            self.parts.append("\t")
        elif ns == _W and name == "noBreakHyphen":
            self.parts.append("-")
        elif ns in (_W, _A) and name in ("br", "cr"):
            self.parts.append("\n")
        elif ns == _A and name == "blip":
            embed = el.get(f"{{{_R}}}embed")
            if embed:
                self.picture_ids.append(embed)
        for child in el:
            self.walk(child)
        if ns in (_W, _A) and name == "p":
            self.parts.append("\n")
        return self


def _xml_text(el: Any) -> str:
    return _clean(_TextWalker().walk(el).text())


def _table_text(rows: Iterator[list[str]]) -> str:
    lines = []
    for cells in rows:
        cells = [re.sub(r"\s+", " ", cell).strip() for cell in cells]
        if any(cells):
            lines.append(" | ".join(cells))
    return "\n".join(lines)


def _m(name: str) -> str:
    return f"{{{_M}}}{name}"


def _simple(text: str) -> bool:
    if re.fullmatch(r"[\w.'′]+", text):
        return True
    if text.startswith("(") and text.endswith(")"):
        depth = 0
        for i, ch in enumerate(text):
            depth += ch == "("
            depth -= ch == ")"
            if depth == 0 and i < len(text) - 1:
                return False
        return True
    return False


def _group(text: str) -> str:
    text = text.strip()
    return text if not text or _simple(text) else f"({text})"


def omml_text(el: Any) -> str:
    """Linear text for an Office Math (OMML) element, like ``dy/dx = cos(x^3+1)*3x^2``.

    Only the structures that matter for reading calculus are spelled out
    (fractions, scripts, roots, delimiters, functions, big operators, limits);
    anything else contributes the text of its parts.
    """
    ns, name = _local(el.tag)
    if ns != _M:
        return "".join(omml_text(child) for child in el) if ns == "" and name == "" else ""

    def part(tag: str) -> str:
        child = el.find(_m(tag))
        return omml_text(child).strip() if child is not None else ""

    def prop(prop_tag: str, tag: str, default: str) -> str:
        child = el.find(f"{_m(prop_tag)}/{_m(tag)}")
        if child is None:
            return default
        return child.get(_m("val"), default)

    if name == "t":
        return el.text or ""
    if name.endswith("Pr"):
        return ""
    if name == "f":
        # Padded so "u" followed by dv/dx reads "u dv/dx", not "udv/dx".
        return f" {_group(part('num'))}/{_group(part('den'))} "
    if name == "sSup":
        return f"{_group(part('e'))}^{_group(part('sup'))}"
    if name == "sSub":
        return f"{_group(part('e'))}_{_group(part('sub'))}"
    if name == "sSubSup":
        return f"{_group(part('e'))}_{_group(part('sub'))}^{_group(part('sup'))}"
    if name == "rad":
        degree = part("deg")
        return f"sqrt({part('e')})" if not degree else f"({part('e')})^(1/{degree})"
    if name == "d":
        begin, end, sep = prop("dPr", "begChr", "("), prop("dPr", "endChr", ")"), prop("dPr", "sepChr", "|")
        inner = sep.join(omml_text(child).strip() for child in el.findall(_m("e")))
        return f"{begin}{inner}{end}"
    if name == "nary":
        symbol = prop("naryPr", "chr", "∫")
        sub, sup = part("sub"), part("sup")
        limits = (f"_{_group(sub)}" if sub else "") + (f"^{_group(sup)}" if sup else "")
        return f"{symbol}{limits} {part('e')}"
    if name == "func":
        arg = part("e")
        return f"{part('fName')}{arg if _simple(arg) and arg.startswith('(') else f'({arg})'}"
    if name == "limLow":
        return f"{part('e')}_{_group(part('lim'))}"
    if name == "limUpp":
        return f"{part('e')}^{_group(part('lim'))}"
    if name == "m":
        rows = [", ".join(omml_text(cell).strip() for cell in row.findall(_m("e"))) for row in el.findall(_m("mr"))]
        return "[" + "; ".join(rows) + "]"
    if name == "eqArr":
        return "; ".join(omml_text(child).strip() for child in el.findall(_m("e")))
    if name in ("acc", "bar", "box", "borderBox", "groupChr", "phant"):
        return part("e")
    if name == "oMathPara":
        return "\n".join(omml_text(child).strip() for child in el.findall(_m("oMath")))
    return "".join(omml_text(child) for child in el)


def _reject_ole(data: bytes) -> None:
    """Encrypted .pptx/.docx files are OLE containers, as are the old .ppt/.doc formats."""
    if data.startswith(_OLE_MAGIC):
        raise ExtractionError(OFFICE_PASSWORD if _ENCRYPTED_PACKAGE in data else OFFICE_LEGACY)


def _check_archive(data: bytes, damaged: str) -> None:
    _reject_ole(data)
    mime, problem = _inspect_zip(data)
    if problem:
        raise ExtractionError(problem)
    if mime is None:
        raise ExtractionError(damaged)


# ---------------------------------------------------------------------------
# PowerPoint


@dataclass
class _SlideItem:
    top: int
    left: int
    text: str


def _extract_pptx(data: bytes, filename: str) -> ExtractedFile:
    from pptx import Presentation

    _check_archive(data, PPTX_DAMAGED)
    try:
        deck = Presentation(io.BytesIO(data))
        slides = list(deck.slides)
        slide_area = int(deck.slide_width or 0) * int(deck.slide_height or 0)
    except Exception as exc:  # python-pptx and lxml raise many kinds of error for broken packages
        raise ExtractionError(PPTX_DAMAGED) from exc
    total = len(slides)
    if total == 0:
        raise ExtractionError(PPTX_EMPTY)

    budget = _ImageBudget()
    pages: list[Page] = []
    warnings: list[str] = []
    for index, slide in enumerate(slides[:MAX_PAGES], start=1):
        try:
            pages.append(_pptx_slide(slide, index, slide_area, budget))
        except Exception:
            pages.append(Page(number=index, label=f"slide {index}"))
            warnings.append(f"Slate couldn't read slide {index}.")

    title = pages[0].heading if pages else ""
    if not title:
        core = _one_line(deck.core_properties.title or "")
        title = core if core.lower() not in {"", "powerpoint presentation"} else _filename_title(filename)
    if total > MAX_PAGES:
        warnings.insert(0, _truncation_warning(total, "slide"))
    result = ExtractedFile(pages=pages, title=title, page_count=total, truncated=total > MAX_PAGES,
                           warnings=warnings)
    return _finish(result, budget)


def _pptx_slide(slide: Any, number: int, slide_area: int, budget: _ImageBudget) -> Page:
    from pptx.shapes.shapetree import SlideShapeFactory

    title_shape = slide.shapes.title
    heading = _one_line(title_shape.text_frame.text) if title_shape is not None else ""
    title_id = title_shape.shape_id if title_shape is not None else None

    items: list[_SlideItem] = []
    pictures: list[tuple[int, Any]] = []
    elements = _slide_shape_elements(slide.shapes._spTree)
    shapes = [SlideShapeFactory(el, slide.shapes) for el in elements]
    _collect_shapes(shapes, items, pictures, skip_id=title_id)
    items.sort(key=lambda item: (item.top, item.left))

    parts = [heading] if heading else []
    parts += [item.text for item in items]
    text = _clean("\n".join(parts))

    picture_area = sum(area for area, _ in pictures)
    share = min(picture_area / slide_area, 1.0) if slide_area else 0.0
    needs_ocr = _needs_ocr(text, share, has_picture=bool(pictures))
    image = None
    if needs_ocr:
        for _area, picture in sorted(pictures, key=lambda pair: pair[0], reverse=True):
            blob = _picture_blob(picture)
            image = _picture_png(blob) if blob else None
            if image is not None:
                break
        needs_ocr = image is not None
        image = budget.take(image)

    if slide.has_notes_slide:
        notes = _clean(slide.notes_slide.notes_text_frame.text) if slide.notes_slide.notes_text_frame else ""
        if notes:
            text = f"{text}\n\nNotes: {notes}".strip()
    return Page(number=number, label=f"slide {number}", heading=heading, text=text, needs_ocr=needs_ocr,
                image=image)


def _picture_blob(picture: Any) -> bytes | None:
    try:
        return picture.image.blob
    except Exception:  # linked (not embedded) pictures have no blob
        return None


def _slide_shape_elements(container: Any) -> list[Any]:
    """Shape elements in a shape tree, including ones PowerPoint wraps in
    ``mc:AlternateContent`` (it does this to any text box holding an equation,
    and python-pptx skips those)."""
    found = []
    for child in container.iterchildren():
        if child.tag == f"{{{_MC}}}AlternateContent":
            choice = child.find(f"{{{_MC}}}Choice")
            if choice is not None:
                found += [el for el in choice.iterchildren() if el.tag in _SLIDE_SHAPE_TAGS]
        elif child.tag in _SLIDE_SHAPE_TAGS:
            found.append(child)
    return found


def _collect_shapes(shapes: list[Any], items: list[_SlideItem], pictures: list[tuple[int, Any]],
                    skip_id: int | None) -> None:
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.shapes.picture import Picture
    from pptx.shapes.shapetree import BaseShapeFactory

    for shape in shapes:
        top, left = int(shape.top or 0), int(shape.left or 0)
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            children = [BaseShapeFactory(el, shape.shapes) for el in _slide_shape_elements(shape._element)]
            group_items: list[_SlideItem] = []
            _collect_shapes(children, group_items, pictures, skip_id=None)
            group_items.sort(key=lambda item: (item.top, item.left))
            text = "\n".join(item.text for item in group_items)
            if text.strip():
                items.append(_SlideItem(top, left, text))
            continue
        if isinstance(shape, Picture):
            area = int(shape.width or 0) * int(shape.height or 0)
            pictures.append((area, shape))
            continue
        if getattr(shape, "has_table", False) and shape.has_table:
            rows = ([_xml_text(cell._tc) for cell in row.cells if not cell.is_spanned] for row in shape.table.rows)
            text = _table_text(rows)
        elif shape.has_text_frame and shape.shape_id != skip_id:
            text = _xml_text(shape._element)
        else:
            continue
        if text.strip():
            items.append(_SlideItem(top, left, text))


# ---------------------------------------------------------------------------
# Word


@dataclass
class _Section:
    heading: str = ""
    lines: list[str] = field(default_factory=list)
    picture_ids: list[str] = field(default_factory=list)
    started_by_heading: bool = False

    @property
    def empty(self) -> bool:
        return not self.lines and not self.picture_ids


def _extract_docx(data: bytes, filename: str) -> ExtractedFile:
    import docx

    _check_archive(data, DOCX_DAMAGED)
    try:
        document = docx.Document(io.BytesIO(data))
        body = document.element.body
        styles = _docx_styles(document.styles.element)
        core_title = _one_line(document.core_properties.title or "")
    except Exception as exc:  # python-docx and lxml raise many kinds of error for broken packages
        raise ExtractionError(DOCX_DAMAGED) from exc

    sections = [_Section()]
    for kind, el in _docx_blocks(body):
        current = sections[-1]
        if kind == "tbl":
            rows = ([_xml_text(tc) for tc in tr.iterchildren(f"{{{_W}}}tc")]
                    for tr in el.iterchildren(f"{{{_W}}}tr"))
            table = _table_text(rows)
            if table:
                current.lines.append(table)
            continue
        walker = _TextWalker().walk(el)
        text = _clean(walker.text())
        role = _paragraph_role(el, styles)
        if role == "heading" and text:
            if not current.empty or current.started_by_heading:
                current = _Section()
                sections.append(current)
            current.heading = _one_line(text)
            current.started_by_heading = True
        elif role == "title" and text and not current.heading:
            current.heading = _one_line(text)
        if text:
            current.lines.append(text)
        current.picture_ids += walker.picture_ids

    sections = [s for s in sections if not s.empty]
    if not sections:
        raise ExtractionError(DOCX_EMPTY)
    total = len(sections)
    budget = _ImageBudget()
    pages = []
    for number, section in enumerate(sections[:MAX_PAGES], start=1):
        text = _clean("\n".join(section.lines))
        needs_ocr = _needs_ocr(text, 0.0, has_picture=bool(section.picture_ids))
        image = _docx_picture(document, section.picture_ids) if needs_ocr else None
        pages.append(Page(number=number, label=f"section {number}", heading=section.heading, text=text,
                          needs_ocr=image is not None, image=budget.take(image)))
    if all(not p.text and not p.needs_ocr for p in pages):
        raise ExtractionError(DOCX_EMPTY)

    result = ExtractedFile(pages=pages, title=core_title or _filename_title(filename), page_count=total,
                           truncated=total > MAX_PAGES,
                           warnings=[_truncation_warning(total, "section")] if total > MAX_PAGES else [])
    return _finish(result, budget)


def _docx_blocks(container: Any) -> Iterator[tuple[str, Any]]:
    """Body paragraphs and tables in order, looking inside content controls."""
    for child in container.iterchildren():
        ns, name = _local(child.tag)
        if ns != _W:
            continue
        if name == "p":
            yield "p", child
        elif name == "tbl":
            yield "tbl", child
        elif name == "sdt":
            content = child.find(f"{{{_W}}}sdtContent")
            if content is not None:
                yield from _docx_blocks(content)


@dataclass
class _StyleInfo:
    name: str
    based_on: str | None
    outline: int | None


def _docx_styles(styles_el: Any) -> dict[str, _StyleInfo]:
    found = {}
    for style in styles_el.iter(f"{{{_W}}}style"):
        style_id = style.get(f"{{{_W}}}styleId")
        if not style_id:
            continue
        name_el = style.find(f"{{{_W}}}name")
        based_el = style.find(f"{{{_W}}}basedOn")
        outline_el = style.find(f"{{{_W}}}pPr/{{{_W}}}outlineLvl")
        found[style_id] = _StyleInfo(
            name=(name_el.get(f"{{{_W}}}val") or "").lower() if name_el is not None else "",
            based_on=based_el.get(f"{{{_W}}}val") if based_el is not None else None,
            outline=_int(outline_el.get(f"{{{_W}}}val")) if outline_el is not None else None,
        )
    return found


def _int(value: str | None) -> int | None:
    try:
        return int(value) if value is not None else None
    except ValueError:
        return None


def _paragraph_role(p: Any, styles: dict[str, _StyleInfo]) -> str:
    """"heading" for Heading styles (or anything with an outline level), "title" for Title, else ""."""
    outline = p.find(f"{{{_W}}}pPr/{{{_W}}}outlineLvl")
    level = _int(outline.get(f"{{{_W}}}val")) if outline is not None else None
    if level is not None and level < 9:
        return "heading"
    style_el = p.find(f"{{{_W}}}pPr/{{{_W}}}pStyle")
    style_id = style_el.get(f"{{{_W}}}val") if style_el is not None else None
    for _ in range(10):  # follow basedOn, which can in theory loop
        info = styles.get(style_id or "")
        if info is None:
            return ""
        if re.match(r"heading \d", info.name) or (info.outline is not None and info.outline < 9):
            return "heading"
        if info.name == "title":
            return "title"
        style_id = info.based_on
    return ""


def _docx_picture(document: Any, picture_ids: list[str]) -> bytes | None:
    blobs = []
    for rid in dict.fromkeys(picture_ids):
        part = document.part.related_parts.get(rid)
        blob = getattr(part, "blob", None)
        if blob:
            blobs.append(blob)
    for blob in sorted(blobs, key=len, reverse=True):
        png = _picture_png(blob)
        if png is not None:
            return png
    return None
